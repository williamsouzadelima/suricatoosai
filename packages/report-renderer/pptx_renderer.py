"""Renderer PPTX (python-pptx) do ReportModel — estilo "dossie".

Casca de dossie: capa com codigo de documento + selo + faixa de KPIs, rodape
numerado (CONFIDENCIAL · docCode · NN), divisorias de PARTE. Marca configuravel
por MSSP (meta.brand sobrescreve a paleta/wordmark). Cada secao de conteudo
vira um slide."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

import theme as T
import logo

EMU_IN = 914400
SW = Inches(13.333)
SH = Inches(7.5)
ML = Inches(0.6)
CW = Inches(12.13)

MONTHS = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
          "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]

# Estado por render (um modelo por processo): meta p/ rodape + contador de pagina.
_STATE = {"meta": {}, "page": 0}


def _rgb(hexstr):
    return RGBColor.from_string(hexstr)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _accent():
    return _STATE["meta"].get("brand", {}).get("accent") or T.CORAL


def _month_year(ms):
    try:
        import datetime
        d = datetime.datetime.utcfromtimestamp(int(ms) / 1000)
        return f"{MONTHS[d.month - 1]} · {d.year}"
    except Exception:
        return ""


def _text(slide, l, t, w, h, text, size=14, bold=False, color=T.INK,
          align=PP_ALIGN.LEFT, font=T.FONT_BODY, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = str(text).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = _rgb(color)
    return tb


def _rect(slide, l, t, w, h, fill, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = _rgb(fill)
    if line:
        sp.line.color.rgb = _rgb(line)
        sp.line.width = Pt(0.75)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _clip(s, n):
    s = str(s or "")
    return s if len(s) <= n else s[: n - 1] + "…"


def _footer(slide, dark=False):
    """Rodape numerado em todo slide de conteudo (cover nao chama)."""
    m = _STATE["meta"]
    _STATE["page"] += 1
    n = _STATE["page"]
    col = "8FA0BC" if dark else T.MUTED
    _text(slide, ML, Inches(7.08), Inches(6.0), Inches(0.3),
          m.get("classification", ""), size=8, color=col)
    _text(slide, Inches(9.0), Inches(7.08), Inches(3.73), Inches(0.3),
          f"{m.get('docCode', '')} · {n:02d}", size=8, color=col,
          align=PP_ALIGN.RIGHT)


def _title_slide(prs, title):
    slide = _blank(prs)
    _rect(slide, 0, 0, SW, Inches(0.12), T.PRIMARY)
    _text(slide, ML, Inches(0.35), CW, Inches(0.7), title, size=24, bold=True,
          color=T.INK, font=T.FONT_DISPLAY)
    _rect(slide, ML, Inches(1.05), Inches(1.2), Emu(int(0.045 * EMU_IN)), _accent())
    _footer(slide)
    return slide, Inches(1.35)


def _cover(prs, sec, meta):
    slide = _blank(prs)
    _rect(slide, 0, 0, SW, SH, T.INK)
    brand = meta.get("brand", {})
    accent = brand.get("accent") or T.CORAL
    # Wordmark (topo-esq): logo embutido so p/ a marca padrao; senao texto.
    used_logo = False
    if brand.get("wordmark", "") == "Suricatoos":
        lp = logo.write_logo(dark=True)
        if lp and os.path.exists(lp):
            try:
                slide.shapes.add_picture(lp, ML, Inches(0.5), height=Inches(0.5))
                used_logo = True
            except Exception:
                used_logo = False
    if not used_logo:
        _text(slide, ML, Inches(0.5), Inches(7.0), Inches(0.5),
              brand.get("wordmark", T.WORDMARK), size=22, bold=True,
              color=accent, font=T.FONT_DISPLAY)
    # Topo-dir: selo + codigo do documento.
    _text(slide, Inches(7.9), Inches(0.5), Inches(4.83), Inches(0.3),
          meta.get("classification", ""), size=10, bold=True, color=accent,
          align=PP_ALIGN.RIGHT)
    _text(slide, Inches(7.9), Inches(0.85), Inches(4.83), Inches(0.3),
          meta.get("docCode", ""), size=10, color="8FA0BC", align=PP_ALIGN.RIGHT)
    # Titulo.
    _rect(slide, ML, Inches(2.2), Inches(1.4), Emu(int(0.06 * EMU_IN)), accent)
    _text(slide, ML, Inches(2.35), CW, Inches(0.35), meta.get("volume", ""),
          size=12, bold=True, color=accent)
    _text(slide, ML, Inches(2.75), CW, Inches(1.3), sec.get("title", ""),
          size=40, bold=True, color="FFFFFF", font=T.FONT_DISPLAY)
    if sec.get("subtitle"):
        _text(slide, ML, Inches(4.1), CW, Inches(0.6), sec["subtitle"], size=18,
              color="C9D3E6")
    # Faixa de KPIs.
    stats = sec.get("stats") or []
    if stats:
        tw = Inches(2.7)
        gap = Inches(0.2)
        yv = Inches(5.15)
        for i, s in enumerate(stats[:4]):
            x = ML + (tw + gap) * i
            _text(slide, x, yv, tw, Inches(0.8), str(s.get("value", "")),
                  size=40, bold=True, color=accent, font=T.FONT_DISPLAY)
            _text(slide, x, yv + Inches(0.85), tw, Inches(0.5),
                  str(s.get("label", "")), size=10, bold=True, color="8FA0BC")
    # Rodape da capa.
    _text(slide, ML, Inches(6.85), CW, Inches(0.4),
          f"{meta.get('client', '')}  ·  {_month_year(meta.get('generatedAt', 0))}  ·  {brand.get('name', '')}",
          size=12, color="C9D3E6")
    _STATE["page"] = 1


def _part_divider(prs, sec):
    slide = _blank(prs)
    _rect(slide, 0, 0, SW, SH, T.INK)
    accent = _accent()
    _rect(slide, ML, Inches(3.0), Inches(1.6), Emu(int(0.06 * EMU_IN)), accent)
    _text(slide, ML, Inches(2.35), CW, Inches(0.4), sec.get("part", ""),
          size=14, bold=True, color=accent)
    _text(slide, ML, Inches(3.2), CW, Inches(1.2), sec.get("title", ""),
          size=34, bold=True, color="FFFFFF", font=T.FONT_DISPLAY)
    if sec.get("subtitle"):
        _text(slide, ML, Inches(4.5), CW, Inches(0.6), sec["subtitle"],
              size=16, color="C9D3E6")
    _footer(slide, dark=True)


def _divider(prs, meta):
    slide = _blank(prs)
    _rect(slide, 0, 0, SW, SH, T.SURFACE_ALT)
    _rect(slide, ML, Inches(3.4), Inches(1.6), Emu(int(0.06 * EMU_IN)), _accent())
    _text(slide, ML, Inches(3.6), CW, Inches(0.6),
          meta.get("brand", {}).get("wordmark", T.WORDMARK), size=22,
          bold=True, color=T.PRIMARY, font=T.FONT_DISPLAY)
    _footer(slide)


def _stat_band(prs, sec):
    slide = _blank(prs)
    _rect(slide, 0, 0, SW, Inches(0.12), T.PRIMARY)
    y = Inches(0.5)
    if sec.get("headline"):
        _text(slide, ML, y, CW, Inches(0.9), _clip(sec["headline"], 160),
              size=24, bold=True, color=T.INK, font=T.FONT_DISPLAY)
        y = Inches(1.55)
    stats = sec.get("stats") or []
    if stats:
        tw = Inches(2.85)
        gap = Inches(0.13)
        for i, s in enumerate(stats[:4]):
            x = ML + (tw + gap) * i
            _rect(slide, x, y, tw, Inches(1.6), T.SURFACE_ALT, line=T.BORDER)
            _rect(slide, x, y, Emu(int(0.06 * EMU_IN)), Inches(1.6), _accent())
            _text(slide, x + Inches(0.2), y + Inches(0.2), tw - Inches(0.3),
                  Inches(0.85), str(s.get("value", "")), size=40, bold=True,
                  color=T.PRIMARY, font=T.FONT_DISPLAY)
            _text(slide, x + Inches(0.2), y + Inches(1.1), tw - Inches(0.3),
                  Inches(0.4), str(s.get("label", "")), size=10, bold=True,
                  color=T.MUTED)
        y = y + Inches(1.95)
    if sec.get("body"):
        _text(slide, ML, y, CW, Inches(2.4), _clip(sec["body"], 720), size=14,
              color=T.INK)
    _footer(slide)


def _attack_chain(prs, sec):
    slide, y = _title_slide(prs, "Cadeia de Ataque")
    lanes = sec.get("lanes") or []
    lane_h = Inches(1.7)
    box_w = Inches(2.0)
    arrow_w = Inches(0.35)
    step = box_w + arrow_w
    for li, lane in enumerate(lanes[:3]):
        ly = y + lane_h * li
        _text(slide, ML, ly, CW, Inches(0.35), _clip(lane.get("title", ""), 80),
              size=12, bold=True, color=T.INK)
        steps = (lane.get("steps") or [])[:5]
        for si, st in enumerate(steps):
            x = ML + step * si
            by = ly + Inches(0.42)
            _rect(slide, x, by, box_w, Inches(0.9), T.SURFACE_ALT, line=T.BORDER)
            _rect(slide, x, by, box_w, Emu(int(0.05 * EMU_IN)), T.PRIMARY)
            _text(slide, x + Inches(0.1), by + Inches(0.1), box_w - Inches(0.2),
                  Inches(0.35), _clip(st.get("label", ""), 22), size=11,
                  bold=True, color=T.PRIMARY)
            if st.get("detail"):
                _text(slide, x + Inches(0.1), by + Inches(0.45),
                      box_w - Inches(0.2), Inches(0.42),
                      _clip(st.get("detail", ""), 44), size=8, color=T.MUTED)
            if si < len(steps) - 1:
                _text(slide, x + box_w, by + Inches(0.28), arrow_w, Inches(0.4),
                      "→", size=18, bold=True, color=_accent(),
                      align=PP_ALIGN.CENTER)


def _roadmap(prs, sec):
    slide, y = _title_slide(prs, "Plano de Remediação por Fase")
    phases = sec.get("phases") or []
    cw = Inches(2.85)
    gap = Inches(0.13)
    for i, ph in enumerate(phases[:4]):
        x = ML + (cw + gap) * i
        _rect(slide, x, y, cw, Inches(5.0), T.SURFACE_ALT, line=T.BORDER)
        _rect(slide, x, y, cw, Inches(0.7), T.PRIMARY)
        _text(slide, x + Inches(0.15), y + Inches(0.08), cw - Inches(0.25),
              Inches(0.3), ph.get("window", ""), size=11, bold=True,
              color="FFFFFF")
        _text(slide, x + Inches(0.15), y + Inches(0.38), cw - Inches(0.25),
              Inches(0.3), ph.get("label", ""), size=10, color="C9D3E6")
        yy = y + Inches(0.9)
        for a in (ph.get("actions") or [])[:6]:
            _text(slide, x + Inches(0.15), yy, cw - Inches(0.3), Inches(0.6),
                  "▸ " + _clip(a, 62), size=9, color=T.INK)
            yy += Inches(0.62)


def _cost_incident(prs, sec):
    slide, y = _title_slide(prs, "Custo da Remediação × Custo do Incidente")
    colw = Inches(5.9)
    gap = Inches(0.33)
    _rect(slide, ML, y, colw, Inches(5.0), T.SURFACE_ALT, line=T.BORDER)
    _rect(slide, ML, y, colw, Inches(0.55), T.SUCCESS)
    _text(slide, ML + Inches(0.2), y + Inches(0.1), colw - Inches(0.3),
          Inches(0.35), "REMEDIAÇÃO — INVESTIMENTO", size=12, bold=True,
          color="FFFFFF")
    yy = y + Inches(0.75)
    for r in (sec.get("remediation") or [])[:6]:
        _text(slide, ML + Inches(0.2), yy, colw - Inches(0.4), Inches(0.5),
              "◆ " + _clip(r, 72), size=11, color=T.INK)
        yy += Inches(0.55)
    x2 = ML + colw + gap
    _rect(slide, x2, y, colw, Inches(5.0), T.SURFACE_ALT, line=T.BORDER)
    _rect(slide, x2, y, colw, Inches(0.55), T.DESTRUCTIVE)
    _text(slide, x2 + Inches(0.2), y + Inches(0.1), colw - Inches(0.3),
          Inches(0.35), "INCIDENTE NÃO REMEDIADO — RISCO", size=12, bold=True,
          color="FFFFFF")
    yy = y + Inches(0.75)
    for it in (sec.get("incident") or [])[:4]:
        _text(slide, x2 + Inches(0.2), yy, colw - Inches(0.4), Inches(0.35),
              "◆ " + _clip(it.get("label", ""), 60), size=12, bold=True,
              color=T.DESTRUCTIVE)
        yy += Inches(0.38)
        if it.get("detail"):
            _text(slide, x2 + Inches(0.45), yy, colw - Inches(0.6), Inches(0.4),
                  _clip(it.get("detail", ""), 72), size=10, color=T.MUTED)
            yy += Inches(0.55)
        else:
            yy += Inches(0.2)


def _partner_value(prs, sec):
    slide, y = _title_slide(prs, _clip(sec.get("title", "Parceria"), 60))
    pts = sec.get("points") or []
    cw = Inches(5.9)
    ch = Inches(2.2)
    gapx = Inches(0.33)
    gapy = Inches(0.3)
    for i, p in enumerate(pts[:4]):
        col = i % 2
        row = i // 2
        x = ML + (cw + gapx) * col
        yy = y + (ch + gapy) * row
        _rect(slide, x, yy, cw, ch, T.SURFACE_ALT, line=T.BORDER)
        _rect(slide, x, yy, Emu(int(0.06 * EMU_IN)), ch, _accent())
        _text(slide, x + Inches(0.25), yy + Inches(0.15), cw - Inches(0.4),
              Inches(0.4), _clip(p.get("title", ""), 40), size=15, bold=True,
              color=T.PRIMARY, font=T.FONT_DISPLAY)
        _text(slide, x + Inches(0.25), yy + Inches(0.6), cw - Inches(0.4),
              Inches(1.5), _clip(p.get("body", ""), 260), size=11, color=T.INK)


def _next_steps(prs, sec):
    slide, y = _title_slide(prs, "Próximos Passos")
    steps = sec.get("steps") or []
    yy = y
    for i, s in enumerate(steps[:6], 1):
        _rect(slide, ML, yy, Inches(0.55), Inches(0.55), T.PRIMARY)
        _text(slide, ML, yy, Inches(0.55), Inches(0.55), str(i), size=20,
              bold=True, color="FFFFFF", align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
        _text(slide, ML + Inches(0.75), yy + Inches(0.05), CW - Inches(1.0),
              Inches(0.5), _clip(s, 110), size=15, color=T.INK,
              anchor=MSO_ANCHOR.MIDDLE)
        yy += Inches(0.75)
    brand = _STATE["meta"].get("brand", {})
    contact = " · ".join(
        [x for x in [brand.get("name", ""), brand.get("contact", ""),
                     brand.get("tagline", "")] if x])
    _text(slide, ML, Inches(6.4), CW, Inches(0.4), contact, size=11,
          color=T.MUTED)


def _kpis(prs, kpis):
    slide, y = _title_slide(prs, "Indicadores")
    tiles = [
        ("Total de achados", str(kpis.get("total", 0)), T.PRIMARY),
        ("Risco (0-100)", str(kpis.get("riskScore", 0)), T.DESTRUCTIVE
         if kpis.get("riskScore", 0) >= 55 else T.WARNING),
        ("CVSS medio", str(kpis.get("avgCvss") if kpis.get("avgCvss") is not None else "-"), T.INK),
        ("Criticos+Altos",
         str(kpis.get("bySeverity", {}).get("critical", 0) + kpis.get("bySeverity", {}).get("high", 0)),
         _accent()),
    ]
    tw = Inches(2.85)
    gap = Inches(0.13)
    for i, (label, val, col) in enumerate(tiles):
        x = ML + (tw + gap) * i
        _rect(slide, x, y, tw, Inches(1.7), T.SURFACE_ALT, line=T.BORDER)
        _rect(slide, x, y, Emu(int(0.06 * EMU_IN)), Inches(1.7), col)
        _text(slide, x + Inches(0.2), y + Inches(0.2), tw - Inches(0.3), Inches(0.4),
              label.upper(), size=10, bold=True, color=T.MUTED)
        _text(slide, x + Inches(0.2), y + Inches(0.6), tw - Inches(0.3), Inches(0.9),
              val, size=40, bold=True, color=col, font=T.FONT_DISPLAY)


def _severity_chart(prs, data):
    slide, y = _title_slide(prs, "Distribuição por severidade")
    dist = data.get("severityDistribution", [])
    maxc = max([d.get("count", 0) for d in dist] + [1])
    barmax = Inches(9.0)
    row_h = Inches(0.62)
    for i, d in enumerate(dist):
        sev = d.get("severity")
        c = d.get("count", 0)
        yy = y + row_h * i
        _text(slide, ML, yy, Inches(2.0), Inches(0.5),
              T.SEVERITY_LABEL.get(sev, sev), size=13, bold=True, color=T.INK,
              anchor=MSO_ANCHOR.MIDDLE)
        w = Emu(int(max(0.04, c / maxc) * barmax)) if c > 0 else Emu(int(0.04 * EMU_IN))
        _rect(slide, ML + Inches(2.2), yy + Inches(0.08), w, Inches(0.42),
              T.SEVERITY_COLOR.get(sev, T.MUTED))
        _text(slide, ML + Inches(2.2) + w + Inches(0.1), yy, Inches(1.0), Inches(0.5),
              str(c), size=13, bold=True, color=T.INK, anchor=MSO_ANCHOR.MIDDLE)


def _risk_matrix(prs, data):
    slide, y = _title_slide(prs, "Matriz de risco (probabilidade × impacto)")
    cells = {(c["likelihood"], c["impact"]): c.get("count", 0)
             for c in data.get("riskMatrix", [])}
    n = 5
    cell = Inches(1.0)
    x0 = ML + Inches(1.0)
    y0 = y + Inches(0.1)

    def band(score):
        if score >= 8:
            return T.DESTRUCTIVE
        if score >= 5:
            return T.WARNING
        if score >= 3:
            return "6F9BF5"
        return T.SUCCESS
    for li in range(n, 0, -1):
        for im in range(1, n + 1):
            row = n - li
            col = im - 1
            x = x0 + cell * col
            yy = y0 + cell * row
            score = li + im
            _rect(slide, x, yy, cell - Inches(0.05), cell - Inches(0.05), band(score))
            cnt = cells.get((li, im), 0)
            if cnt:
                _text(slide, x, yy, cell - Inches(0.05), cell - Inches(0.05),
                      str(cnt), size=20, bold=True, color="FFFFFF",
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _text(slide, x0, y0 + cell * n + Inches(0.05), cell * n, Inches(0.4),
          "Probabilidade →", size=11, color=T.MUTED)


def _finding_card(prs, sec):
    f = sec["finding"]
    sev = f.get("severity", "info")
    slide, y = _title_slide(prs, _clip(f.get("title", "Achado"), 70))
    _rect(slide, ML, y, CW, Inches(0.5), T.SEVERITY_COLOR.get(sev, T.MUTED))
    _text(slide, ML + Inches(0.15), y + Inches(0.05), CW, Inches(0.4),
          f"{T.SEVERITY_LABEL.get(sev, sev).upper()}  ·  {f.get('ref','')}  ·  {_clip(f.get('affectedAsset',''),60)}",
          size=13, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
    yy = y + Inches(0.7)
    meta = []
    if f.get("weaknessClass"):
        meta.append(f"Classe: {f['weaknessClass']}")
    if f.get("cvssVector"):
        meta.append(f"CVSS: {f['cvssVector']}")
    if f.get("cwe"):
        meta.append(f"CWE: {f['cwe']}")
    if meta:
        _text(slide, ML, yy, CW, Inches(0.4), "   ".join(meta), size=11, color=T.MUTED)
        yy += Inches(0.45)
    for label, key in [("Descrição", "description"), ("Impacto", "impact"),
                       ("Remediação", "remediation")]:
        if f.get(key):
            _text(slide, ML, yy, CW, Inches(0.3), label.upper(), size=10, bold=True, color=T.PRIMARY)
            yy += Inches(0.28)
            _text(slide, ML, yy, CW, Inches(0.8), _clip(f[key], 500), size=12, color=T.INK)
            yy += Inches(0.85)
    if sec.get("showPoc") and f.get("reproductionSteps"):
        steps = "\n".join(f"{i+1}. {_clip(s,140)}" for i, s in enumerate(f["reproductionSteps"][:6]))
        _text(slide, ML, yy, CW, Inches(0.3), "REPRODUÇÃO", size=10, bold=True, color=T.PRIMARY)
        _text(slide, ML, yy + Inches(0.28), CW, Inches(1.4), steps, size=11, color=T.INK)
    # Evidencia (texto + imagem) em slides dedicados — so em relatorios com PoC.
    if sec.get("showPoc"):
        _evidence_slides(prs, f)


def _image_slide(prs, title, path, caption=None):
    if not path or not os.path.exists(path):
        return
    slide, y = _title_slide(prs, title)
    try:
        pic = slide.shapes.add_picture(path, ML, y)
    except Exception:
        return
    max_w = int(CW)
    max_h = int(Inches(5.0))
    scale = min(max_w / pic.width, max_h / pic.height, 1.0)
    pic.width = int(pic.width * scale)
    pic.height = int(pic.height * scale)
    pic.left = int(ML + (int(CW) - pic.width) / 2)
    pic.top = int(y)
    if caption:
        _text(slide, ML, int(y) + pic.height + Inches(0.1), CW, Inches(0.4),
              _clip(caption, 120), size=11, color=T.MUTED)


def _evidence_slides(prs, f):
    ev = sorted(
        f.get("evidence") or [],
        key=lambda e: e.get("stepIndex") if e.get("stepIndex") is not None else 1_000_000,
    )
    chain = [e for e in ev if e.get("snippet") or e.get("command")]
    imgs = [e for e in ev if e.get("imagePath")]

    if f.get("narrative"):
        slide, y = _title_slide(prs, _clip(f"Narrativa — {f.get('title','')}", 70))
        _text(slide, ML, y, CW, Inches(5.0), _clip(f["narrative"], 1400),
              size=14, color=T.INK)

    if chain:
        slide, y = _title_slide(
            prs, _clip(f"Cadeia de evidência — {f.get('title','')}", 70))
        tb = slide.shapes.add_textbox(ML, y, CW, Inches(5.2))
        tf = tb.text_frame
        tf.word_wrap = True
        first = True

        def _line(text, size, bold, color, before=0):
            nonlocal first
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = T.FONT_BODY
            r.font.color.rgb = _rgb(color)
            if before:
                p.space_before = Pt(before)

        for i, e in enumerate(chain[:6], 1):
            n = e.get("stepIndex") or i
            head = f"{n}."
            if e.get("toolName"):
                head += f"  [{e['toolName']}]"
            _line(head, 12, True, T.PRIMARY, before=6)
            if e.get("command"):
                _line("$ " + _clip(e["command"], 180), 10, False, T.MUTED)
            if e.get("snippet"):
                _line(_clip(e["snippet"], 340), 10, False, T.INK)
            if e.get("resultSummary"):
                _line("-> prova: " + _clip(e["resultSummary"], 180), 10, False,
                      T.SUCCESS)

    for e in imgs[:4]:
        _image_slide(prs, _clip(f"Evidência Visual — {f.get('title','')}", 70),
                     e.get("imagePath"), e.get("label"))


def _table(prs, findings, title, remediation=False):
    slide, y = _title_slide(prs, title)
    rows = min(len(findings), 12) + 1
    if remediation:
        headers = ["Ref", "Sev", "Ativo", "Remediacao"]
        widths = [Inches(1.2), Inches(1.4), Inches(3.5), Inches(6.0)]
    else:
        headers = ["Ref", "Sev", "Titulo", "Ativo"]
        widths = [Inches(1.2), Inches(1.4), Inches(5.5), Inches(4.0)]
    gt = slide.shapes.add_table(rows, len(headers), ML, y, CW, Inches(0.4) * rows).table
    for j, hd in enumerate(headers):
        gt.columns[j].width = widths[j]
        cell = gt.cell(0, j)
        cell.text = hd
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(T.PRIMARY)
        pr = cell.text_frame.paragraphs[0].runs[0]
        pr.font.size = Pt(11)
        pr.font.bold = True
        pr.font.color.rgb = _rgb("FFFFFF")
    for i, f in enumerate(findings[:12], start=1):
        sev = f.get("severity", "info")
        vals = ([f.get("ref", ""), T.SEVERITY_LABEL.get(sev, sev), _clip(f.get("affectedAsset", ""), 40), _clip(f.get("remediation", ""), 90)]
                if remediation else
                [f.get("ref", ""), T.SEVERITY_LABEL.get(sev, sev), _clip(f.get("title", ""), 70), _clip(f.get("affectedAsset", ""), 40)])
        for j, v in enumerate(vals):
            cell = gt.cell(i, j)
            cell.text = str(v)
            r = cell.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(10)
            r.font.color.rgb = _rgb(T.SEVERITY_COLOR.get(sev, T.INK) if j == 1 else T.INK)
            if j == 1:
                r.font.bold = True


def _bullets(prs, items, title):
    slide, y = _title_slide(prs, title or "Itens")
    tb = slide.shapes.add_textbox(ML, y, CW, Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items[:16]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + _clip(it, 130)
        r.font.size = Pt(14)
        r.font.name = T.FONT_BODY
        r.font.color.rgb = _rgb(T.INK)
        p.space_after = Pt(6)


def _paragraph(prs, text, title):
    slide, y = _title_slide(prs, title or "")
    _text(slide, ML, y, CW, Inches(5.0), _clip(text, 1200), size=15, color=T.INK)


def _callout(prs, sec, title):
    slide, y = _title_slide(prs, title or "")
    col = T.CALLOUT_COLOR.get(sec.get("tone", "info"), T.PRIMARY)
    _rect(slide, ML, y, CW, Inches(1.2), T.SURFACE_ALT, line=T.BORDER)
    _rect(slide, ML, y, Emu(int(0.08 * EMU_IN)), Inches(1.2), col)
    _text(slide, ML + Inches(0.3), y + Inches(0.15), CW - Inches(0.5), Inches(0.9),
          _clip(sec.get("text", ""), 400), size=16, bold=True, color=col,
          anchor=MSO_ANCHOR.MIDDLE)


def _apply_brand(meta):
    brand = meta.get("brand") or {}
    if brand.get("primary"):
        T.PRIMARY = brand["primary"]
    if brand.get("accent"):
        T.CORAL = brand["accent"]
    if brand.get("wordmark"):
        T.WORDMARK = brand["wordmark"]


def render_pptx(model, out_path):
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    meta = model.get("meta", {})
    _STATE["meta"] = meta
    _STATE["page"] = 0
    _apply_brand(meta)
    title = None
    for sec in model.get("sections", []):
        t = sec.get("type")
        if t == "cover":
            _cover(prs, sec, meta)
        elif t == "partDivider":
            _part_divider(prs, sec)
        elif t == "divider":
            _divider(prs, meta)
        elif t == "heading":
            title = sec.get("text")
        elif t == "statBand":
            _stat_band(prs, sec)
        elif t == "kpis":
            _kpis(prs, sec.get("kpis", {}))
        elif t == "severityChart":
            _severity_chart(prs, sec.get("data", {}))
        elif t == "riskMatrix":
            _risk_matrix(prs, sec.get("data", {}))
        elif t == "attackChain":
            _attack_chain(prs, sec)
        elif t == "findingCard":
            _finding_card(prs, sec)
        elif t == "findingsTable":
            _table(prs, sec.get("findings", []), title or "Sumário de Achados")
        elif t == "remediationMatrix":
            _table(prs, sec.get("findings", []), title or "Matriz de Remediação", remediation=True)
        elif t == "roadmap":
            _roadmap(prs, sec)
        elif t == "costVsIncident":
            _cost_incident(prs, sec)
        elif t == "partnerValue":
            _partner_value(prs, sec)
        elif t == "nextSteps":
            _next_steps(prs, sec)
        elif t == "bullets":
            _bullets(prs, sec.get("items", []), title)
        elif t == "paragraph":
            _paragraph(prs, sec.get("text", ""), title)
        elif t == "callout":
            _callout(prs, sec, title)
    prs.save(out_path)
