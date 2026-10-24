"""Renderer DOCX (python-docx) do ReportModel — estilo "dossie".

Casca de dossie: capa com codigo de documento + faixa de KPIs, rodape com
classificacao + docCode + numero de pagina, divisorias de PARTE. Marca
configuravel por MSSP (meta.brand sobrescreve paleta/wordmark)."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
import os
import re

import theme as T
import logo


def _rgb(h):
    return RGBColor.from_string(h)


def _shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def _cell_borders(cell, specs):
    """Bordas da célula numa única w:tcBorders. specs: {edge: (color_hex, sz)},
    sz em oitavos de ponto (4≈0.5pt, 24=3pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, (color, sz) in specs.items():
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


_EV_FILE_RE = re.compile(
    r'^\s*[\w.\-/]+\.(?:png|jpe?g|gif|webp|bmp|svg|tiff?)\s*(?:[—–:\-]\s*)?', re.I)
_STEP_NUM_RE = re.compile(r'^\s*\d{1,3}[.)\]]\s+')


def _clean_ev_label(label, toolname=None):
    """Rótulo de evidência legível ao cliente: remove o nome de ARQUIVO cru
    ('radar.png — ...'); se era só o nome do arquivo, retorna ''."""
    s = str(label or "").strip()
    if not s:
        return ""
    s2 = _EV_FILE_RE.sub("", s).strip()
    if not s2:
        return ""
    if toolname and s2 == str(toolname):
        return ""
    return s2


def _strip_step_num(s):
    """Remove numeração que o passo já traz ('1. ...'); o renderer renumera."""
    return _STEP_NUM_RE.sub("", str(s or "").strip())


def _set_font(r, name):
    """Define a familia do run em ascii+hAnsi+cs — garante IBM Plex em todo o
    texto (ASCII, acentos pt-BR e scripts complexos), nao so no ASCII."""
    r.font.name = name
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), name)


def _run(p, text, size=11, bold=False, color=T.INK, font=T.FONT_BODY):
    r = p.add_run(str(text))
    r.font.size = Pt(size)
    r.font.bold = bold
    _set_font(r, font)
    r.font.color.rgb = _rgb(color)
    return r


def _para(doc, text, size=11, bold=False, color=T.INK, after=6, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    _run(p, text, size, bold, color)
    return p


def _heading(doc, text, level=1):
    p = doc.add_paragraph()
    # Estilo de cabeçalho nomeado → entra no outline/navegação/TOC do Word.
    # A formatação da marca (run abaixo) sobrepõe a aparência do estilo.
    try:
        p.style = doc.styles["Heading 1" if level == 1 else "Heading 2"]
    except KeyError:
        pass
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    _run(p, text, size=17 if level == 1 else 14, bold=True, color=T.PRIMARY,
         font=T.FONT_DISPLAY)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), T.CORAL)
    bdr.append(bottom)
    pPr.append(bdr)
    return p


def _clip(s, n):
    s = str(s or "")
    return s if len(s) <= n else s[: n - 1] + "…"


# Salvaguarda contra entradas patológicas; NÃO trunca conteúdo real de achado.
# DOCX pagina sozinho — corpo de achado/evidência nunca é cortado (regra do dossiê).
BODY_SOFT_MAX = 200000


def _body(doc, text, size=11, bold=False, color=T.INK, after=6, font=None,
          align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Parágrafo de corpo que preserva quebras de linha (\\n) e NUNCA trunca.
    Justificado por padrão (casa com PDF/PPTX) — só o corpo de prosa passa por
    aqui (descrição/impacto/remediação/narrativa/cadeia/passos); rótulos, títulos
    e legendas usam _para e não são afetados."""
    s = str(text or "")
    if len(s) > BODY_SOFT_MAX:
        s = s[:BODY_SOFT_MAX]
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run()
    r.font.size = Pt(size)
    r.font.bold = bold
    _set_font(r, font or T.FONT_BODY)
    r.font.color.rgb = _rgb(color)
    parts = s.split("\n")
    for i, ln in enumerate(parts):
        if i:
            r.add_break()
        r.add_text(ln)
    return p


def _apply_brand(meta):
    brand = meta.get("brand") or {}
    if brand.get("primary"):
        T.PRIMARY = brand["primary"]
    if brand.get("accent"):
        T.CORAL = brand["accent"]
    if brand.get("wordmark"):
        T.WORDMARK = brand["wordmark"]


def _add_page_field(p, size=8, color=None):
    color = color or T.MUTED
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")  # sem isto o nº de página cai p/ serifa
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), T.FONT_BODY)
    rpr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    rpr.append(sz)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    rpr.append(col)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)


def _setup_footer(doc, meta):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.text = ""
    _run(p, meta.get("classification", ""), size=8, color=T.MUTED)
    _run(p, "   ·   " + meta.get("docCode", "") + "   ·   pág. ", size=8,
         color=T.MUTED)
    _add_page_field(p)


def _stat_tiles(doc, stats):
    stats = (stats or [])[:4]
    if not stats:
        return
    t = doc.add_table(rows=2, cols=len(stats))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, s in enumerate(stats):
        c0 = t.cell(0, j)
        c0.text = ""
        _shade(c0, T.SURFACE_ALT)
        _run(c0.paragraphs[0], str(s.get("value", "")), size=26, bold=True,
             color=T.PRIMARY, font=T.FONT_DISPLAY)
        c1 = t.cell(1, j)
        c1.text = ""
        _shade(c1, T.SURFACE_ALT)
        _run(c1.paragraphs[0], str(s.get("label", "")).upper(), size=9,
             bold=True, color=T.MUTED)
    doc.add_paragraph()


# Namespaces OOXML p/ a âncora de imagem flutuante (capa full-bleed).
_WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"


def _page_background(doc, img_path, pw_emu, ph_emu):
    """Imagem full-page ancorada à PÁGINA, ATRÁS do texto (behindDoc + wrapNone).
    No Word, relativeFrom=page + offset 0 = canto físico da página = full-bleed
    real (independe das margens da seção); o parágrafo-portador fica na pág. 1, então
    a arte só aparece na capa. OBS: o LibreOffice trata page como área de texto e
    mostra margem branca — o alvo é o Word (entregável do cliente)."""
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(img_path, width=Emu(int(pw_emu)), height=Emu(int(ph_emu)))
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'<wp:anchor xmlns:wp="{_WP_NS}" xmlns:a="{_A_NS}" xmlns:r="{_R_NS}" '
        f'xmlns:pic="{_PIC_NS}" behindDoc="1" distT="0" distB="0" distL="0" '
        'distR="0" simplePos="0" locked="0" layoutInCell="1" allowOverlap="1" '
        'relativeHeight="0">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset>'
        "</wp:positionH>"
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset>'
        "</wp:positionV>"
        f'<wp:extent cx="{int(pw_emu)}" cy="{int(ph_emu)}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        "<wp:wrapNone/>"
        '<wp:docPr id="1000" name="CoverBackground"/>'
        "<wp:cNvGraphicFramePr/>"
        "</wp:anchor>"
    )
    anchor.append(graphic)
    drawing.remove(inline)
    drawing.append(anchor)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    return p


def _cover_stats_white(doc, stats):
    """Faixa de KPIs da capa full-bleed: texto branco, SEM sombreamento (células
    transparentes deixam o navy do fundo aparecer)."""
    stats = (stats or [])[:4]
    if not stats:
        return
    t = doc.add_table(rows=2, cols=len(stats))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j, s in enumerate(stats):
        c0 = t.cell(0, j)
        c0.text = ""
        _run(c0.paragraphs[0], str(s.get("value", "")), size=24, bold=True,
             color="FFFFFF", font=T.FONT_DISPLAY)
        c1 = t.cell(1, j)
        c1.text = ""
        _run(c1.paragraphs[0], str(s.get("label", "")).upper(), size=9,
             bold=True, color="C9D3E6")
    doc.add_paragraph()


def _cover(doc, sec, meta):
    brand = meta.get("brand", {})
    # Capa com arte da marca (coverImagePath) → FULL-BLEED real: a arte preenche a
    # página inteira (imagem flutuante behindDoc) e o texto branco fica por cima,
    # como no PPTX/PDF. Ausente → capa programática clara (título escuro).
    cover = brand.get("coverImagePath")
    if not (cover and os.path.exists(cover)):
        # Sem arte da marca → a capa NASCE navy (padrão-ouro SYN), não clara:
        # fundo navy full-page + título/logo/KPIs em branco.
        try:
            cover = T.default_cover_bg()
        except Exception:
            cover = None
    fullbleed = bool(cover and os.path.exists(cover))
    if fullbleed:
        secpg = doc.sections[0]
        try:
            fitted = T.cover_fit(cover, int(secpg.page_width),
                                 int(secpg.page_height))
            _page_background(doc, fitted, secpg.page_width, secpg.page_height)
        except Exception:
            fullbleed = False
    # Wordmark (topo-esq): logo custom > logo Suricatoos (claro p/ navy full-bleed,
    # escuro p/ capa clara) > texto.
    used_logo = False
    logo_path = brand.get("logoPath")
    if logo_path and os.path.exists(logo_path):
        try:
            doc.add_picture(logo_path, width=Inches(1.1))
            used_logo = True
        except Exception:
            used_logo = False
    if not used_logo and brand.get("wordmark", "") == "Suricatoos":
        lp = logo.write_logo(dark=fullbleed)
        if lp and os.path.exists(lp):
            try:
                # Wordmark oficial Suricatoos (1600x274, ~5.84:1) — destaque de
                # capa; largura generosa p/ o logotipo dominar o topo-esquerdo.
                doc.add_picture(lp, width=Inches(2.6))
                used_logo = True
            except Exception:
                used_logo = False
    if not used_logo:
        _para(doc, brand.get("wordmark", T.WORDMARK), size=22, bold=True,
              color=T.CORAL, align=WD_ALIGN_PARAGRAPH.LEFT)
    title_color = "FFFFFF" if fullbleed else T.INK
    muted_color = "C9D3E6" if fullbleed else T.MUTED
    if fullbleed:
        for _ in range(6):
            doc.add_paragraph()
    _para(doc, f"{meta.get('classification', '')}   ·   {meta.get('docCode', '')}",
          size=10, bold=True, color=T.CORAL, after=2)
    _para(doc, meta.get("volume", ""), size=10, bold=True, color=muted_color,
          after=10)
    _para(doc, sec.get("title", ""), size=34, bold=True, color=title_color,
          after=4)
    if sec.get("subtitle"):
        _para(doc, sec["subtitle"], size=16, color=muted_color, after=10)
    if fullbleed:
        _cover_stats_white(doc, sec.get("stats"))
    else:
        _stat_tiles(doc, sec.get("stats"))
    doc.add_page_break()


def _part_divider(doc, sec, lead_break=True):
    # Sem quebra de entrada quando vem logo após a capa (que já quebrou) — evita a
    # página em branco entre a capa e a 1ª PARTE. Divisores no meio do documento
    # mantêm a quebra (começam em página nova).
    if lead_break:
        doc.add_page_break()
    for _ in range(3):
        doc.add_paragraph()
    _para(doc, sec.get("part", ""), size=13, bold=True, color=T.CORAL, after=2)
    _outline_title(doc, sec.get("title", ""), 30, T.INK, level=1, after=4)
    if sec.get("subtitle"):
        _para(doc, sec["subtitle"], size=15, color=T.MUTED)


def _stat_band(doc, sec):
    if sec.get("headline"):
        _para(doc, sec["headline"], size=20, bold=True, color=T.INK,
              after=8)
    _stat_tiles(doc, sec.get("stats"))
    if sec.get("body"):
        _body(doc, sec["body"], size=11, color=T.INK, after=6)


def _attack_chain(doc, sec):
    _heading(doc, "Cadeia de Ataque", 1)
    for lane in (sec.get("lanes") or []):
        _para(doc, lane.get("title", ""), size=12, bold=True,
              color=T.INK, after=2)
        steps = (lane.get("steps") or [])
        if steps:
            _para(doc, "  →  ".join(s.get("label", "") for s in steps),
                  size=11, bold=True, color=T.PRIMARY, after=2)
        for s in steps:
            if s.get("detail"):
                p = doc.add_paragraph(style="List Bullet")
                _run(p, f"{s.get('label','')}: {s['detail']}",
                     size=10, color=T.MUTED)
        doc.add_paragraph()


def _roadmap(doc, sec):
    _heading(doc, "Plano de Remediação por Fase", 1)
    phases = (sec.get("phases") or [])[:4]
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for j, hd in enumerate(["Janela", "Foco", "Ações"]):
        c = t.cell(0, j)
        c.text = ""
        _shade(c, T.PRIMARY)
        _run(c.paragraphs[0], hd, size=10, bold=True, color="FFFFFF")
    for ph in phases:
        row = t.add_row()
        _run(row.cells[0].paragraphs[0], ph.get("window", ""), size=10,
             bold=True, color=T.INK)
        _run(row.cells[1].paragraphs[0], ph.get("label", ""), size=10,
             color=T.MUTED)
        acts = ph.get("actions") or []
        cell = row.cells[2]
        cell.text = ""
        for k, a in enumerate(acts):
            p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            _run(p, "▸ " + str(a), size=9, color=T.INK)
    doc.add_paragraph()


def _cost_incident(doc, sec):
    _heading(doc, "Custo da Remediação × Custo do Incidente", 1)
    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    h0 = t.cell(0, 0)
    h0.text = ""
    _shade(h0, T.SUCCESS)
    _run(h0.paragraphs[0], "REMEDIAÇÃO — INVESTIMENTO", size=10, bold=True,
         color="FFFFFF")
    h1 = t.cell(0, 1)
    h1.text = ""
    _shade(h1, T.DESTRUCTIVE)
    _run(h1.paragraphs[0], "INCIDENTE NÃO REMEDIADO — RISCO", size=10, bold=True,
         color="FFFFFF")
    cr = t.cell(1, 0)
    cr.text = ""
    for k, r in enumerate((sec.get("remediation") or [])):
        p = cr.paragraphs[0] if k == 0 else cr.add_paragraph()
        _run(p, "◆ " + str(r), size=10, color=T.INK)
    ci = t.cell(1, 1)
    ci.text = ""
    inc = (sec.get("incident") or [])
    for k, it in enumerate(inc):
        p = ci.paragraphs[0] if k == 0 else ci.add_paragraph()
        _run(p, "◆ " + it.get("label", ""), size=10, bold=True,
             color=T.DESTRUCTIVE)
        if it.get("detail"):
            pp = ci.add_paragraph()
            _run(pp, "   " + str(it["detail"]), size=9, color=T.MUTED)
    doc.add_paragraph()


def _partner_value(doc, sec):
    _heading(doc, sec.get("title", "Parceria"), 1)
    for p in (sec.get("points") or [])[:6]:
        _para(doc, p.get("title", ""), size=13, bold=True,
              color=T.PRIMARY, after=1)
        _para(doc, p.get("body", ""), size=11, color=T.INK, after=6)


def _next_steps(doc, sec):
    _heading(doc, "Próximos Passos", 1)
    for i, s in enumerate((sec.get("steps") or [])[:6], 1):
        _para(doc, f"{i}.  {s}", size=12, color=T.INK, after=4)
    brand = sec.get("_brand") or {}
    contact = " · ".join(
        [x for x in [brand.get("name", ""), brand.get("contact", ""),
                     brand.get("tagline", "")] if x])
    if contact:
        _para(doc, contact, size=10, color=T.MUTED, after=2)


def _kpis(doc, kpis):
    tiles = [
        ("Total de achados", str(kpis.get("total", 0))),
        ("Risco (0-100)", str(kpis.get("riskScore", 0))),
        ("CVSS médio", str(kpis.get("avgCvss") if kpis.get("avgCvss") is not None else "-")),
        ("Críticos+Altos", str(kpis.get("bySeverity", {}).get("critical", 0)
                               + kpis.get("bySeverity", {}).get("high", 0))),
    ]
    t = doc.add_table(rows=2, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (label, val) in enumerate(tiles):
        c0 = t.cell(0, j)
        c0.text = ""
        _shade(c0, T.SURFACE_ALT)
        _run(c0.paragraphs[0], label.upper(), size=9, bold=True, color=T.MUTED)
        c1 = t.cell(1, j)
        c1.text = ""
        _shade(c1, T.SURFACE_ALT)
        _run(c1.paragraphs[0], val, size=26, bold=True, color=T.PRIMARY,
             font=T.FONT_DISPLAY)
    doc.add_paragraph()


def _severity_chart(doc, data):
    # Barras horizontais de verdade: célula sombreada com a cor da severidade e
    # largura proporcional à contagem (antes era "████" em texto). A célula
    # restante fica sem borda/preenchimento (invisível) — só a barra aparece.
    dist = data.get("severityDistribution", [])
    maxc = max([d.get("count", 0) for d in dist] + [1])
    TRACK = 5.6  # polegadas da trilha total da barra
    for d in dist:
        sev = d.get("severity")
        c = d.get("count", 0)
        color = T.SEVERITY_COLOR.get(sev, T.MUTED)
        lbl = doc.add_paragraph()
        lbl.paragraph_format.space_before = Pt(4)
        lbl.paragraph_format.space_after = Pt(1)
        _run(lbl, f"{T.SEVERITY_LABEL.get(sev, sev)}  ·  {c}", size=10,
             bold=True, color=color)
        frac = (c / maxc) if maxc else 0.0
        barw = max(0.06, TRACK * frac)
        t = doc.add_table(rows=1, cols=2)
        t.allow_autofit = False
        bar, rest = t.cell(0, 0), t.cell(0, 1)
        bar.width = Inches(barw)
        rest.width = Inches(max(0.02, TRACK - barw))
        _shade(bar, color)
        for cell in (bar, rest):
            pc = cell.paragraphs[0]
            pc.paragraph_format.space_before = Pt(0)
            pc.paragraph_format.space_after = Pt(0)
            r = pc.add_run(" ")
            r.font.size = Pt(7)
    doc.add_paragraph()


def _risk_matrix(doc, data):
    cells = {(c["likelihood"], c["impact"]): c.get("count", 0)
             for c in data.get("riskMatrix", [])}

    def band(score):
        if score >= 8:
            return T.DESTRUCTIVE
        if score >= 5:
            return T.WARNING
        if score >= 3:
            return "6F9BF5"
        return T.SUCCESS
    t = doc.add_table(rows=5, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in range(5):
        li = 5 - r
        for col in range(5):
            im = col + 1
            cell = t.cell(r, col)
            cell.text = ""
            _shade(cell, band(li + im))
            cnt = cells.get((li, im), 0)
            _run(cell.paragraphs[0], str(cnt) if cnt else "", size=12,
                 bold=True, color="FFFFFF")
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def _finding_card(doc, sec):
    f = sec["finding"]
    sev = f.get("severity", "info")
    p = doc.add_paragraph()
    # Título do achado como Heading 3: fica no outline/navegação do Word MAS fora
    # do campo TOC (\o "1-2" pega só níveis 1-2) → "atualizar campo" no Word não
    # inunda o SUMÁRIO com os achados. A aparência vem dos runs abaixo.
    try:
        p.style = doc.styles["Heading 3"]
    except KeyError:
        pass
    p.paragraph_format.space_before = Pt(10)
    _run(p, f"[{T.SEVERITY_LABEL.get(sev, sev).upper()}] ", size=13, bold=True,
         color=T.SEVERITY_COLOR.get(sev, T.MUTED))
    _run(p, f"{f.get('ref','')} · {f.get('title','')}", size=13,
         bold=True, color=T.INK)
    meta = []
    if f.get("affectedAsset"):
        meta.append(f"Ativo: {f['affectedAsset']}")
    if f.get("weaknessClass"):
        meta.append(f"Classe: {f['weaknessClass']}")
    if f.get("cvssVector"):
        meta.append(f"CVSS: {f['cvssVector']}")
    if f.get("cwe"):
        meta.append(f"CWE: {f['cwe']}")
    if meta:
        _para(doc, "   ".join(meta), size=10, color=T.MUTED, after=4)
    if f.get("narrative"):
        _para(doc, "NARRATIVA:", size=10, bold=True, color=T.PRIMARY, after=1)
        _body(doc, f["narrative"], size=10, color=T.INK, after=4)
    for label, key in [("Descrição", "description"), ("Impacto", "impact"),
                       ("Remediação", "remediation")]:
        if f.get(key):
            _para(doc, label.upper() + ":", size=10, bold=True,
                  color=T.PRIMARY, after=1)
            _body(doc, f[key], size=11, color=T.INK, after=6)
    if sec.get("showPoc") and f.get("reproductionSteps"):
        _para(doc, "REPRODUÇÃO:", size=10, bold=True, color=T.PRIMARY, after=2)
        for i, s in enumerate(f["reproductionSteps"], 1):
            _body(doc, f"{i}. {_strip_step_num(s)}", size=10, color=T.INK, after=1)
    if sec.get("showPoc"):
        _evidence_docx(doc, f)


def _embed_image_docx(doc, path, caption=None):
    if not path or not os.path.exists(path):
        return False
    try:
        # Print EMOLDURADO: célula de 1x1 com borda fina (exibit), em vez de uma
        # imagem solta flutuando no branco (tímida).
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        cell.text = ""
        _shade(cell, T.SURFACE)
        _cell_borders(cell, {e: (T.BORDER, 4)
                             for e in ("top", "left", "bottom", "right")})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(path, width=Inches(5.8))
    except Exception:
        return False
    cap = _clean_ev_label(caption)
    cap = ("◆ " + cap) if cap else "◆ Evidência visual"
    _para(doc, cap, size=8, color=T.MUTED, after=8,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    return True


def _placeholder_docx(doc, text, color=None):
    """Nota NEUTRA de evidência complementar (sem alarme). Só p/ achado de fato
    não-verificável (sem cadeia textual E sem imagem) — nunca sob achado que já
    traz prova técnica."""
    color = color or T.MUTED
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    cell.text = ""
    _shade(cell, T.SURFACE_ALT)
    _run(cell.paragraphs[0], text, size=10, bold=True, color=color)
    doc.add_paragraph()


def _ev_sorted(f):
    ev = f.get("evidence") or []
    return sorted(
        ev,
        key=lambda e: e.get("stepIndex") if e.get("stepIndex") is not None else 1_000_000,
    )


# Cores da cadeia de evidência em painel CLARO (fundo SURFACE_ALT). As antigas
# EV_* da marca são claras (feitas p/ fundo escuro de terminal) e sumiriam aqui.
_EVL_HEAD = T.PRIMARY      # cabeçalho do passo (azul, destaque)
_EVL_LABEL = T.MUTED       # rótulo/descrição do passo
_EVL_CMD = T.INK           # comando (escuro, "como digitado")
_EVL_OUT = "334155"        # output/snippet (ardósia escura, legível)
_EVL_OK = T.SUCCESS        # prompt "$" e "→ prova" (verde)


def _evidence_block(doc, steps):
    """Cadeia de evidência em PAINEL CLARO ("exhibit"): célula de fundo tingido
    (SURFACE_ALT) com barra de acento azul à esquerda e texto MONO de ALTO
    contraste — cabeçalho azul, comando escuro, output ardósia, prova verde. Um
    bloco por achado (toda a cadeia junta); a célula do Word quebra entre páginas."""
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    cell.text = ""
    _shade(cell, T.SURFACE_ALT)
    _cell_borders(cell, {"top": (T.BORDER, 4), "bottom": (T.BORDER, 4),
                         "right": (T.BORDER, 4), "left": (T.PRIMARY, 24)})
    state = {"first": True}

    def line(before=0):
        p = cell.paragraphs[0] if state["first"] else cell.add_paragraph()
        state["first"] = False
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(0)
        return p

    def mono(p, text, color, bold=False):
        r = p.add_run()
        r.font.size = Pt(8.5)
        r.font.bold = bold
        _set_font(r, T.FONT_MONO)
        r.font.color.rgb = _rgb(color)
        for i, ln in enumerate(str(text).split("\n")):
            if i:
                r.add_break()
            r.add_text(ln)

    for idx, (n, e) in enumerate(steps):
        ph = line(before=(2 if idx == 0 else 8))
        head = f"{n}."
        if e.get("toolName"):
            head += f" [{e['toolName']}]"
        mono(ph, head, _EVL_HEAD, bold=True)
        lbl = _clean_ev_label(e.get("label"), e.get("toolName"))
        if lbl:
            mono(ph, "  " + lbl, _EVL_LABEL)
        if e.get("command"):
            pc = line()
            mono(pc, "$ ", _EVL_OK, bold=True)
            mono(pc, str(e["command"]), _EVL_CMD)
        if e.get("snippet"):
            mono(line(), e["snippet"], _EVL_OUT)
        if e.get("resultSummary"):
            pr = line()
            mono(pr, "→ prova: ", _EVL_OK, bold=True)
            mono(pr, e["resultSummary"], T.INK)
    doc.add_paragraph()


def _evidence_docx(doc, f):
    ev = _ev_sorted(f)
    steps = [(e.get("stepIndex") or i, e) for i, e in enumerate(ev, 1)]
    text_steps = [
        (n, e) for (n, e) in steps
        if e.get("command") or e.get("snippet") or e.get("resultSummary")
    ]
    img_steps = [e for (n, e) in steps if e.get("imagePath")]
    if text_steps:
        _para(doc, "CADEIA DE EVIDÊNCIA", size=9, bold=True, color=T.PRIMARY,
              after=3)
        _evidence_block(doc, text_steps)
    img_rendered = 0
    # Só anuncia imagem que REALMENTE existe no pacote; print ainda não enviado
    # (relatório defasado) cai fora SEM aviso — a cadeia textual sustenta.
    existing = [e for e in img_steps
                if e.get("imagePath") and os.path.exists(e.get("imagePath"))]
    for e in existing:
        if _embed_image_docx(doc, e.get("imagePath"), e.get("label")):
            img_rendered += 1
    # Nota neutra APENAS p/ achado de fato não-verificável: sem imagem E sem
    # cadeia técnica. Dizer "não capturada"/"removida" sob achado com prova
    # textual seria FALSO; "indisponível" é sinal interno, não vai ao cliente.
    if (img_rendered == 0 and not text_steps
            and f.get("severity") in ("critical", "high")):
        _placeholder_docx(doc, "Evidência complementar não anexada a este achado.")


def _fixed_table(doc, headers, widths):
    """Tabela de layout FIXO (não deixa o autofit esmagar colunas), cabeçalho
    navy que REPETE em cada página, larguras por coluna (polegadas)."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.allow_autofit = False
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(lay)
    hr = t.rows[0]
    th = OxmlElement("w:tblHeader")  # repete o cabeçalho a cada página
    th.set(qn("w:val"), "true")
    hr._tr.get_or_add_trPr().append(th)
    for j, hd in enumerate(headers):
        c = t.cell(0, j)
        c.text = ""
        c.width = Inches(widths[j])
        _shade(c, T.PRIMARY)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        _run(p, hd, size=9, bold=True, color="FFFFFF")
    return t


def _table(doc, findings, remediation=False):
    if remediation:
        headers = ["Ref", "Sev", "Ativo", "Remediação"]
        widths = [0.75, 0.85, 1.95, 3.35]
        clips = [None, None, 70, 240]
    else:
        headers = ["Ref", "Sev", "Título", "Ativo", "CVSS"]
        widths = [0.75, 0.85, 3.05, 1.75, 0.55]
        clips = [None, None, 130, 58, None]
    t = _fixed_table(doc, headers, widths)
    for idx, f in enumerate(findings):
        sev = f.get("severity", "info")
        row = t.add_row()
        fill = T.SURFACE_ALT if idx % 2 == 0 else T.SURFACE  # zebra
        cvss = f.get("cvssScore")
        cvss_s = f"{cvss:.1f}" if isinstance(cvss, (int, float)) else "—"
        raw = ([f.get("ref", ""), T.SEVERITY_LABEL.get(sev, sev),
                f.get("affectedAsset", ""), f.get("remediation", "")]
               if remediation else
               [f.get("ref", ""), T.SEVERITY_LABEL.get(sev, sev),
                f.get("title", ""), f.get("affectedAsset", ""), cvss_s])
        for j, v in enumerate(raw):
            cell = row.cells[j]
            cell.text = ""
            cell.width = Inches(widths[j])
            _shade(cell, fill)
            val = v  # sem truncar: a célula do Word quebra linha (nada cortado)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            _run(p, val, size=9, bold=(j == 1),
                 color=(T.SEVERITY_COLOR.get(sev, T.INK) if j == 1 else T.INK))
    doc.add_paragraph()


def _callout(doc, sec):
    col = T.CALLOUT_COLOR.get(sec.get("tone", "info"), T.PRIMARY)
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    cell.text = ""
    _shade(cell, T.SURFACE_ALT)
    _run(cell.paragraphs[0], _clip(sec.get("text", ""), 2000), size=12, bold=True,
         color=col)
    doc.add_paragraph()


def _outline_title(doc, text, size, color, level=1, after=4, before=0):
    """Título com estilo Heading NOMEADO (entra no SUMÁRIO/outline do Word) mas
    com a aparência da marca via run — sem a borda do _heading."""
    p = doc.add_paragraph()
    try:
        p.style = doc.styles["Heading 1" if level == 1 else "Heading 2"]
    except KeyError:
        pass
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    _run(p, text, size=size, bold=True, color=color, font=T.FONT_DISPLAY)
    return p


def _toc(doc, entries):
    """SUMÁRIO à prova de falha: campo TOC (níveis 1-2: partes/seções/fases;
    títulos de achado são Heading 3 → ficam de fora) marcado `dirty` → o Word
    popula as páginas ao atualizar. O RESULTADO do campo já traz a LISTA ESTÁTICA
    das seções, então, atualizando ou não, NUNCA fica vazio/placeholder (vira o
    sumário paginado OU a lista de seções limpa)."""
    _para(doc, "SUMÁRIO", size=20, bold=True, color=T.INK, after=10)

    def _fldchar(p, ftype, dirty=False):
        r = p.add_run()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), ftype)
        if dirty:
            fc.set(qn("w:dirty"), "true")
        r._r.append(fc)

    p0 = doc.add_paragraph()
    _fldchar(p0, "begin", dirty=True)
    r = p0.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r._r.append(it)
    _fldchar(p0, "separate")
    items = [e for e in (entries or []) if e and e[0]]
    if not items:
        _fldchar(p0, "end")
        doc.add_page_break()
        return
    for i, (title, level) in enumerate(items):
        pe = doc.add_paragraph()
        pe.paragraph_format.space_after = Pt(5)
        if level:
            pe.paragraph_format.left_indent = Inches(0.28 * level)  # recuo hierárquico
        _run(pe, title, size=(12 if level == 0 else 11),
             bold=(level == 0), color=T.INK)
        if i == len(items) - 1:
            _fldchar(pe, "end")  # fecha o campo na última entrada
    doc.add_page_break()


def render_docx(model, out_path):
    doc = Document()
    meta = model.get("meta", {})
    _apply_brand(meta)
    style = doc.styles["Normal"]
    style.font.name = T.FONT_BODY
    style.font.size = Pt(11)
    _setup_footer(doc, meta)
    brand = meta.get("brand", {})
    # SUMÁRIO completo: partes (nível 0) + headings de seção (nível 1) + as
    # sub-seções/fases da narrativa (nível 2). Títulos de achado são findingCard
    # (não "heading"), então NÃO entram — o sumário não incha com os N achados.
    toc_entries = []
    for s in model.get("sections", []):
        st = s.get("type")
        if st == "partDivider":
            toc_entries.append((s.get("title", ""), 0))
        elif st == "heading":
            lvl = s.get("level", 1)
            if lvl in (1, 2):
                toc_entries.append((s.get("text", ""), lvl))
    prev_type = None
    for sec in model.get("sections", []):
        t = sec.get("type")
        if t == "cover":
            _cover(doc, sec, meta)
            _toc(doc, toc_entries)
        elif t == "partDivider":
            _part_divider(doc, sec, lead_break=(prev_type != "cover"))
        elif t == "divider":
            doc.add_page_break()
        elif t == "heading":
            _heading(doc, sec.get("text", ""), sec.get("level", 1))
        elif t == "paragraph":
            _body(doc, sec.get("text", ""), size=11, color=T.INK)
        elif t == "bullets":
            for it in sec.get("items", []):
                p = doc.add_paragraph(style="List Bullet")
                _run(p, str(it), size=11, color=T.INK)
        elif t == "statBand":
            _stat_band(doc, sec)
        elif t == "kpis":
            _kpis(doc, sec.get("kpis", {}))
        elif t == "severityChart":
            _severity_chart(doc, sec.get("data", {}))
        elif t == "riskMatrix":
            _risk_matrix(doc, sec.get("data", {}))
        elif t == "attackChain":
            _attack_chain(doc, sec)
        elif t == "findingCard":
            _finding_card(doc, sec)
        elif t == "findingsTable":
            _table(doc, sec.get("findings", []))
        elif t == "remediationMatrix":
            _table(doc, sec.get("findings", []), remediation=True)
        elif t == "roadmap":
            _roadmap(doc, sec)
        elif t == "costVsIncident":
            _cost_incident(doc, sec)
        elif t == "partnerValue":
            _partner_value(doc, sec)
        elif t == "nextSteps":
            sec["_brand"] = brand
            _next_steps(doc, sec)
        elif t == "callout":
            _callout(doc, sec)
        prev_type = t
    doc.save(out_path)
    # Embute IBM Plex no pacote OOXML (best-effort: falha → DOCX válido só com o
    # nome da fonte). Sem isso o Word do cliente cai p/ serifa de substituição.
    try:
        import fontembed
        fontembed.embed_fonts(out_path)
    except Exception:
        pass
